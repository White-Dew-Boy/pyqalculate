"""Tests for XLSX import/export functionality.

Tests:
- Calculator.importXLSX() — loads .xlsx as matrix or per-column variables
- Calculator.exportXLSX() — writes MathStructure to .xlsx
"""
import os

import pytest

from pyqalculate.calculator import Calculator
from pyqalculate.math_structure import MathStructure


@pytest.fixture
def calc():
    c = Calculator()
    c.load_definitions()
    return c


@pytest.fixture
def sample_xlsx(tmp_path):
    """Create a sample .xlsx with mixed data."""
    import openpyxl
    xlsx_file = tmp_path / "test_data.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name", "Score", "Grade"])
    ws.append(["Alice", 95, "A"])
    ws.append(["Bob", 87, "B"])
    ws.append(["Charlie", 92, "A"])
    wb.save(str(xlsx_file))
    return str(xlsx_file)


@pytest.fixture
def numeric_xlsx(tmp_path):
    """Create .xlsx with only numeric data."""
    import openpyxl
    xlsx_file = tmp_path / "numbers.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([1.5, 2.5, 3.5])
    ws.append([4.5, 5.5, 6.5])
    ws.append([7.5, 8.5, 9.5])
    wb.save(str(xlsx_file))
    return str(xlsx_file)


@pytest.fixture
def multi_type_xlsx(tmp_path):
    """Create .xlsx with int, float, bool, None, str cells."""
    import openpyxl
    xlsx_file = tmp_path / "mixed.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Label", "Int", "Float", "Bool", "Empty"])
    ws.append(["row1", 42, 3.14, True, None])
    ws.append(["row2", 0, -1.5, False, ""])
    wb.save(str(xlsx_file))
    return str(xlsx_file)


@pytest.fixture
def multi_sheet_xlsx(tmp_path):
    """Create .xlsx with two worksheets."""
    import openpyxl
    xlsx_file = tmp_path / "multi.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "First"
    ws1.append([1, 2, 3])
    ws2 = wb.create_sheet("Second")
    ws2.append([4, 5, 6])
    wb.save(str(xlsx_file))
    return str(xlsx_file)


class TestImportXLSX:
    """Test Calculator.importXLSX()."""

    def test_import_as_matrix(self, calc, numeric_xlsx):
        result = calc.importXLSX(numeric_xlsx, headers=False, to_matrix=True, name="mymatrix")
        assert result.is_matrix()
        assert len(result) == 3
        for row in result:
            assert row.is_vector()
            assert len(row) == 3

    def test_import_as_vectors(self, calc, sample_xlsx):
        result = calc.importXLSX(sample_xlsx, headers=True, to_matrix=False, name="data")
        assert result.is_matrix()
        assert len(result) == 3

    def test_import_registers_matrix_variable(self, calc, numeric_xlsx):
        calc.importXLSX(numeric_xlsx, to_matrix=True, name="mydata")
        assert calc.get_variable("mydata") is not None

    def test_import_registers_column_variables(self, calc, sample_xlsx):
        calc.importXLSX(sample_xlsx, headers=True, to_matrix=False, name="scores")
        assert calc.get_variable("scores_Name") is not None
        assert calc.get_variable("scores_Score") is not None
        assert calc.get_variable("scores_Grade") is not None

    def test_import_skip_first_rows(self, calc, numeric_xlsx):
        result = calc.importXLSX(numeric_xlsx, first_row=2, headers=False, to_matrix=True)
        assert result.is_matrix()
        assert len(result) == 2

    def test_import_by_sheet_index(self, calc, multi_sheet_xlsx):
        result = calc.importXLSX(multi_sheet_xlsx, sheet=1, headers=False, to_matrix=True, name="second")
        assert result.is_matrix()

    def test_import_by_sheet_name(self, calc, multi_sheet_xlsx):
        result = calc.importXLSX(multi_sheet_xlsx, sheet="Second", headers=False, to_matrix=True, name="named")
        assert result.is_matrix()

    def test_import_numeric_values(self, calc, numeric_xlsx):
        result = calc.importXLSX(numeric_xlsx, headers=False, to_matrix=True)
        first_cell = result[0][0]
        assert first_cell.is_number()
        assert abs(first_cell.float_value() - 1.5) < 1e-10

    def test_import_mixed_types(self, calc, multi_type_xlsx):
        result = calc.importXLSX(multi_type_xlsx, headers=True, to_matrix=True, name="mixed")
        assert result.is_matrix()
        row1 = result[0]
        int_cell = row1[1]  # "Int" column
        assert int_cell.is_number()
        assert abs(int_cell.float_value() - 42.0) < 1e-10

    def test_import_bool_as_0_1(self, calc, multi_type_xlsx):
        result = calc.importXLSX(multi_type_xlsx, headers=True, to_matrix=True, name="mixed")
        row1 = result[0]
        bool_cell = row1[3]  # "Bool" column: True → "1"
        assert bool_cell.is_number()
        assert abs(bool_cell.float_value() - 1.0) < 1e-10

    def test_import_none_as_zero(self, calc, multi_type_xlsx):
        result = calc.importXLSX(multi_type_xlsx, headers=True, to_matrix=True, name="mixed")
        row1 = result[0]
        empty_cell = row1[4]  # "Empty" column: None → 0
        assert empty_cell.is_number()
        assert abs(empty_cell.float_value() - 0.0) < 1e-10

    def test_import_nonexistent_returns_undefined(self, calc):
        result = calc.importXLSX("/nonexistent/path/file.xlsx")
        assert result.is_undefined()

    def test_import_empty_file_returns_undefined(self, calc, tmp_path):
        import openpyxl
        empty_file = tmp_path / "empty.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "Empty"
        wb.save(str(empty_file))
        result = calc.importXLSX(str(empty_file))
        assert result.is_undefined()

    def test_import_auto_name_from_filename(self, calc, numeric_xlsx):
        calc.importXLSX(numeric_xlsx, to_matrix=True)
        var = calc.get_variable("numbers")
        assert var is not None

    def test_import_string_values_become_symbols(self, calc, sample_xlsx):
        calc.importXLSX(sample_xlsx, headers=True, to_matrix=False, name="test")
        name_var = calc.get_variable("test_Name")
        assert name_var is not None


class TestExportXLSX:
    """Test Calculator.exportXLSX()."""

    def test_export_matrix(self, calc, tmp_path):
        row1 = MathStructure.vector(MathStructure(1), MathStructure(2), MathStructure(3))
        row2 = MathStructure.vector(MathStructure(4), MathStructure(5), MathStructure(6))
        matrix = MathStructure.matrix([row1, row2])
        out_file = str(tmp_path / "output.xlsx")
        success = calc.exportXLSX(matrix, out_file)
        assert success is True
        import openpyxl
        wb = openpyxl.load_workbook(out_file)
        rows = list(wb.active.iter_rows(values_only=True))
        assert len(rows) == 2
        for i, expected_row in enumerate(([1, 2, 3], [4, 5, 6])):
            for j, expected in enumerate(expected_row):
                assert rows[i][j] == expected

    def test_export_vector_as_column(self, calc, tmp_path):
        vec = MathStructure.vector(MathStructure(10), MathStructure(20), MathStructure(30))
        out_file = str(tmp_path / "vector.xlsx")
        success = calc.exportXLSX(vec, out_file)
        assert success is True
        import openpyxl
        wb = openpyxl.load_workbook(out_file)
        rows = list(wb.active.iter_rows(values_only=True))
        assert len(rows) == 3
        assert rows[0][0] == 10
        assert rows[1][0] == 20
        assert rows[2][0] == 30

    def test_export_single_value(self, calc, tmp_path):
        val = MathStructure(42)
        out_file = str(tmp_path / "single.xlsx")
        success = calc.exportXLSX(val, out_file)
        assert success is True
        import openpyxl
        wb = openpyxl.load_workbook(out_file)
        rows = list(wb.active.iter_rows(values_only=True))
        assert len(rows) == 1
        assert rows[0][0] == 42

    def test_roundtrip_import_export_import(self, calc, numeric_xlsx, tmp_path):
        result1 = calc.importXLSX(numeric_xlsx, headers=False, to_matrix=True, name="rt")
        out_file = str(tmp_path / "roundtrip.xlsx")
        calc.exportXLSX(result1, out_file)
        calc2 = Calculator()
        calc2.load_definitions()
        result2 = calc2.importXLSX(out_file, headers=False, to_matrix=True, name="rt2")
        assert result2.is_matrix()
        assert len(result1) == len(result2)
        for i in range(len(result1)):
            for j in range(len(result1[i])):
                v1 = result1[i][j].float_value()
                v2 = result2[i][j].float_value()
                assert abs(v1 - v2) < 1e-10

    def test_export_custom_sheet_name(self, calc, tmp_path):
        vec = MathStructure.vector(MathStructure(1), MathStructure(2))
        out_file = str(tmp_path / "custom_sheet.xlsx")
        calc.exportXLSX(vec, out_file, sheet_name="MyData")
        import openpyxl
        wb = openpyxl.load_workbook(out_file)
        assert wb.active.title == "MyData"

    def test_export_invalid_path_returns_false(self, calc):
        val = MathStructure(1)
        success = calc.exportXLSX(val, "/nonexistent/dir/file.xlsx")
        assert success is False
